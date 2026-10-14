import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from zipfile import BadZipFile
from openpyxl import load_workbook


COLUNAS_GNRE = (
    "Código Domínio",
    "UF Favorecida",
    "Mês",
    "Ano",
    "IE",
    "Receita",
    "Data Vencimento",
    "Valor Tipo 11",
)

UFS_VALIDAS = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO",
    "MA", "MT", "MS", "MG", "PA", "PB", "PR", "PE", "PI",
    "RJ", "RN", "RS", "RO", "RR", "SC", "SP", "SE", "TO",
}

MAX_LINHAS = 5000
MAX_LINHAS_NA_RESPOSTA = 200


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _inteiro(valor, campo: str, minimo: int, maximo: int) -> int:
    texto = _texto(valor)
    if not texto:
        raise ValueError(f"{campo} é obrigatório.")

    if isinstance(valor, bool):
        raise ValueError(f"{campo} deve ser um número inteiro.")

    try:
        numero = Decimal(texto)
    except InvalidOperation:
        raise ValueError(f"{campo} deve ser um número inteiro.") from None

    if not numero.is_finite() or numero != numero.to_integral_value():
        raise ValueError(f"{campo} deve ser um número inteiro.")

    numero = int(numero)
    if not minimo <= numero <= maximo:
        raise ValueError(f"{campo} deve estar entre {minimo} e {maximo}.")
    return numero


def _valor_monetario(valor) -> Decimal | None:
    if valor is None or _texto(valor) == "":
        return None

    if isinstance(valor, bool) or isinstance(valor, (date, datetime)):
        raise ValueError("Valor Tipo 11 inválido.")

    if isinstance(valor, (int, float, Decimal)):
        texto = str(valor)
    else:
        texto = _texto(valor).replace("R$", "").replace(" ", "").replace("\xa0", "")
        if "," in texto and "." in texto:
            if texto.rfind(",") > texto.rfind("."):
                texto = texto.replace(".", "").replace(",", ".")
            else:
                texto = texto.replace(",", "")
        elif "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        elif texto.count(".") > 1:
            texto = texto.replace(".", "")
        elif re.fullmatch(r"[+-]?\d{1,3}\.\d{3}", texto):
            # Texto no formato brasileiro "1.234" (milhar).
            texto = texto.replace(".", "")

    try:
        quantia = Decimal(texto)
    except InvalidOperation:
        raise ValueError("Valor Tipo 11 deve ser um valor numérico.") from None

    if not quantia.is_finite():
        raise ValueError("Valor Tipo 11 inválido.")

    return quantia.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _proximo_mes(ano: int, mes: int, dia: int) -> date:
    if mes == 12:
        ano, mes = ano + 1, 1
    else:
        mes += 1

    try:
        return date(ano, mes, dia)
    except ValueError:
        raise ValueError(
            f"O dia {dia} não existe no mês seguinte à competência "
            f"({mes:02d}/{ano})."
        ) from None


def _validar_linha(linha_excel: int, valores: dict) -> dict:
    """
    Retorna o resultado da validação de uma linha da planilha.
    Não executa emissões e não utiliza banco de dados.
    """
    dados = {"linha_excel": linha_excel}
    erros: list[str] = []
    avisos: list[str] = []

    try:
        valor = _valor_monetario(valores["Valor Tipo 11"])
    except ValueError as exc:
        valor = None
        erros.append(str(exc))

    # Somente valor não preenchido é ignorado.
    # Valores zerados (0,00) são válidos e devem ser processados.
    if not erros and valor is None:
        return {
            "linha_excel": linha_excel,
            "status": "ignorada",
            "erros": [],
            "avisos": ["Linha ignorada: Valor Tipo 11 não preenchido."],
            "dados": {},
        }

    if valor is not None:
        if valor < 0:
            erros.append("Valor Tipo 11 não pode ser negativo.")
        if valor > Decimal("99999999.99"):
            erros.append("Valor Tipo 11 excede o limite de 10 dígitos do banco.")
        dados["valor_tipo_11"] = f"{valor:.2f}"

    dominio = _texto(valores["Código Domínio"])
    if not dominio:
        erros.append("Código Domínio é obrigatório.")
    elif len(dominio) > 100:
        erros.append("Código Domínio excede 100 caracteres.")
    dados["codigo_dominio"] = dominio

    uf = _texto(valores["UF Favorecida"]).upper()
    if uf not in UFS_VALIDAS:
        erros.append("UF Favorecida deve ser uma sigla de UF brasileira válida.")
    dados["uf_favorecida"] = uf

    mes = ano = dia = None
    for campo, destino, minimo, maximo in (
        ("Mês", "mes", 1, 12),
        ("Ano", "ano", 2000, 9998),
        ("Data Vencimento", "dia_vencimento", 1, 31),
    ):
        try:
            numero = _inteiro(valores[campo], campo, minimo, maximo)
            dados[destino] = numero
            if destino == "mes":
                mes = numero
            if destino == "ano":
                ano = numero
            if destino == "dia_vencimento":
                dia = numero
        except ValueError as exc:
            erros.append(str(exc))

    if all(n is not None for n in (mes, ano, dia)):
        try:
            dados["data_vencimento_calculada"] = (
                _proximo_mes(ano, mes, dia).isoformat()
            )
        except ValueError as exc:
            erros.append(str(exc))

    ie = _texto(valores["IE"])
    if not ie:
        erros.append("IE é obrigatória.")
    elif len(ie) > 50:
        erros.append("IE excede 50 caracteres.")
    elif isinstance(valores["IE"], (int, float)):
        avisos.append(
            "IE veio como número: confirme se zeros iniciais não foram perdidos."
        )
    dados["inscricao_estadual"] = ie

    receita = _texto(valores["Receita"])
    if not receita:
        erros.append("Receita é obrigatória.")
    elif len(receita) > 50:
        erros.append("Receita excede 50 caracteres.")
    elif isinstance(valores["Receita"], (int, float)):
        avisos.append(
            "Receita veio como número: confirme se zeros iniciais não foram perdidos."
        )
    dados["receita"] = receita

    return {
        "linha_excel": linha_excel,
        "status": "erro" if erros else "valida",
        "erros": erros,
        "avisos": avisos,
        "dados": dados,
    }


def analisar_planilha(conteudo: bytes) -> dict:
    """
    Lê e valida o Excel GNRE sem persistir dados ou enviar SOAP.
    Retorna resumo completo e prévia das primeiras 200 linhas úteis.
    """
    try:
        pasta = load_workbook(
            io.BytesIO(conteudo),
            read_only=True,
            data_only=True,
        )
    except (BadZipFile, OSError, ValueError, TypeError, KeyError) as exc:
        raise ValueError("O arquivo não é um Excel .xlsx válido.") from exc

    try:
        aba = pasta["GNRE"] if "GNRE" in pasta.sheetnames else pasta.active

        linhas = aba.iter_rows(values_only=True)
        cabecalhos = next(linhas, None)
        if cabecalhos is None:
            raise ValueError("A planilha não possui cabeçalho.")

        colunas = [
            _texto(c).removeprefix("\ufeff")
            for c in cabecalhos
        ]

        duplicadas = sorted({
            c for c in COLUNAS_GNRE if colunas.count(c) > 1
        })
        if duplicadas:
            raise ValueError(
                "Colunas duplicadas: " + ", ".join(duplicadas)
            )

        ausentes = [c for c in COLUNAS_GNRE if c not in colunas]
        if ausentes:
            raise ValueError(
                "Colunas obrigatórias ausentes: " + ", ".join(ausentes)
            )

        indices = {coluna: colunas.index(coluna) for coluna in COLUNAS_GNRE}

        qtd_validas = qtd_ignoradas = qtd_erros = total = 0
        previa: list[dict] = []

        for numero_excel, linha in enumerate(linhas, start=2):
            valores = {
                nome: linha[indice] if indice < len(linha) else None
                for nome, indice in indices.items()
            }

            if all(_texto(valor) == "" for valor in valores.values()):
                continue

            total += 1
            if total > MAX_LINHAS:
                raise ValueError(
                    f"Limite de {MAX_LINHAS} linhas preenchidas por arquivo excedido."
                )

            resultado = _validar_linha(numero_excel, valores)
            status = resultado["status"]

            if status == "valida":
                qtd_validas += 1
            elif status == "ignorada":
                qtd_ignoradas += 1
            else:
                qtd_erros += 1

            if len(previa) < MAX_LINHAS_NA_RESPOSTA:
                previa.append(resultado)

        return {
            "total_linhas": total,
            "validas": qtd_validas,
            "ignoradas": qtd_ignoradas,
            "com_erro": qtd_erros,
            "pronta_para_confirmacao": qtd_validas > 0 and qtd_erros == 0,
            "previa_limitada": total > MAX_LINHAS_NA_RESPOSTA,
            "linhas": previa,
            "observacao": (
                "Validação de formato e campos, sem emissão de GNRE. "
                "A existência/autorização do certificado, as regras da UF "
                "e a aceitação pela GNRE só serão verificadas na etapa de emissão."
            ),
        }
    finally:
        pasta.close()