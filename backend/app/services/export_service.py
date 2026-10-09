from __future__ import annotations
import tempfile
from datetime import date, datetime
from decimal import Decimal
from typing import BinaryIO
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session
from app.models import GNREProcessamento


MAX_REGISTROS_EXCEL = 40000
CABECALHOS = [
    "UF Favorecida",
    "Código Domínio",
    "IE",
    "Receita",
    "Mês",
    "Ano",
    "Data Vencimento",
    "Valor Tipo 11",
    "Recibo",
    "OBS",
    "Status",
    "ID",
    "Data Envio",
    "PDF Armazenado",
]
LARGURAS = [17, 20, 23, 17, 10, 11, 20, 20, 25, 58, 22, 12, 23, 20]


def filtros_guias(
    *,
    codigo_dominio: str | None = None,
    mes: int | None = None,
    ano: int | None = None,
    status: str | None = None,
    uf_favorecida: str | None = None,
    ids: list[int] | None = None,
) -> list:
    filtros = []

    if codigo_dominio:
        filtros.append(GNREProcessamento.codigo_dominio == codigo_dominio)
    if mes is not None:
        filtros.append(GNREProcessamento.mes == mes)
    if ano is not None:
        filtros.append(GNREProcessamento.ano == ano)
    if status:
        filtros.append(GNREProcessamento.status == status)
    if uf_favorecida:
        filtros.append(GNREProcessamento.uf_favorecida == uf_favorecida.upper())
    if ids is not None:
        filtros.append(GNREProcessamento.id.in_(ids))

    return filtros


def tratar_texto_excel(valor):
    """Evita que conteúdo de empresa ou mensagens seja interpretado como fórmula."""
    if valor is None:
        return ""
    if isinstance(valor, (date, datetime, Decimal, int, float)):
        return valor
    texto = str(valor)
    # Inclusão de apóstrofo antes de expressões potencialmente executáveis.
    if texto.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + texto
    return texto


def gerar_excel_guias(db: Session, *, filtros: list) -> tuple[BinaryIO, int]:
    """Gera planilha XLSX formatada em arquivo temporário."""
    total = db.scalar(
        select(func.count(GNREProcessamento.id)).where(*filtros)
    ) or 0

    if total > MAX_REGISTROS_EXCEL:
        raise ValueError(
            f"A exportação contém {total} guias. "
            f"Use filtros para exportar até {MAX_REGISTROS_EXCEL} por vez."
        )

    consulta = select(
        GNREProcessamento.uf_favorecida,
        GNREProcessamento.codigo_dominio,
        GNREProcessamento.inscricao_estadual,
        GNREProcessamento.receita,
        GNREProcessamento.mes,
        GNREProcessamento.ano,
        GNREProcessamento.data_vencimento,
        GNREProcessamento.valor_tipo_11,
        GNREProcessamento.numero_recibo,
        GNREProcessamento.erro_mensagem,
        GNREProcessamento.status,
        GNREProcessamento.id,
        GNREProcessamento.data_envio,
        case(
            (func.length(GNREProcessamento.pdf_base64) > 0, "Sim"),
            else_="Não",
        ).label("pdf_armazenado"),
    ).where(*filtros).order_by(GNREProcessamento.id.desc())

    arquivo = tempfile.SpooledTemporaryFile(max_size=16 * 1024 * 1024, mode="w+b")
    try:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Resultados")
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for indice, largura in enumerate(LARGURAS, start=1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(indice)].width = largura

        cabecalho = []
        for titulo in CABECALHOS:
            celula = WriteOnlyCell(ws, value=titulo)
            celula.fill = PatternFill("solid", fgColor="15324A")
            celula.font = Font(color="FFFFFF", bold=True)
            celula.alignment = Alignment(vertical="center")
            cabecalho.append(celula)
        ws.append(cabecalho)

        for valores in db.execute(consulta).yield_per(500):
            linha = []
            tem_erro = bool(valores.erro_mensagem)
            for indice, valor in enumerate(valores, start=1):
                celula = WriteOnlyCell(ws, value=tratar_texto_excel(valor))

                if indice == 7 and valor is not None:
                    celula.number_format = "DD/MM/YYYY"
                elif indice == 8 and valor is not None:
                    celula.number_format = '#,##0.00'
                elif indice == 13 and valor is not None:
                    celula.number_format = "DD/MM/YYYY HH:MM"
                elif indice in (2, 3, 4, 9):
                    celula.number_format = "@"

                if indice == 10 and tem_erro:
                    celula.fill = PatternFill("solid", fgColor="FEE2E2")
                    celula.font = Font(color="991B1B")
                    celula.alignment = Alignment(wrap_text=True)

                linha.append(celula)

            ws.append(linha)

        wb.save(arquivo)
        arquivo.seek(0)
        return arquivo, total
    except Exception:
        arquivo.close()
        raise
